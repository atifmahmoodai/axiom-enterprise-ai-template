"""Trusted, bounded analytical programs. No user expressions or generated code are evaluated."""
from __future__ import annotations
import hashlib
import json
import math
import os
import re
import sys
import warnings
import zipfile
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd

MAX_ROWS = 1_000_000
MAX_COLUMNS = 128
MAX_DECODED_BYTES = 200 * 1024 * 1024


def clean(value):
    if value is None:
        return None
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, np.ndarray)):
        return [clean(v) for v in value]
    try:
        if pd.isna(value):
            return None
    except (ValueError, TypeError):
        pass
    return value


def validate_frame(frame):
    if len(frame) > MAX_ROWS:
        raise ValueError("Decoded input exceeds one million rows.")
    if len(frame.columns) > MAX_COLUMNS:
        raise ValueError("A dataset may contain at most 128 columns.")
    if not len(frame.columns):
        raise ValueError("The file does not contain a table header.")
    columns = [str(c).strip() for c in frame.columns]
    if any(not c or len(c) > 120 or "\x00" in c for c in columns):
        raise ValueError("Column names must be nonempty and at most 120 characters.")
    if len(set(columns)) != len(columns):
        raise ValueError("Column names must be unique.")
    frame.columns = columns
    if int(frame.memory_usage(index=True, deep=True).sum()) > MAX_DECODED_BYTES:
        raise ValueError("Decoded dataset exceeds the 200 MiB memory budget.")
    for column in columns:
        if frame[column].map(lambda v: isinstance(v, (dict, list, bytes))).any():
            raise ValueError("Nested objects, arrays and binary cells require flattening before upload.")
        if frame[column].map(lambda v: isinstance(v, str) and len(v) > 65536).any():
            raise ValueError("A cell exceeds the 64 KiB text limit.")
    return frame


def read_table(file_path, extension):
    file_path = Path(file_path)
    if file_path.stat().st_size > 50 * 1024 * 1024 and extension != "canonical":
        raise ValueError("Upload exceeds 50 MiB.")
    if extension == "csv":
        # Header inspection prevents pandas from silently suffixing duplicate labels.
        import csv
        with open(file_path, newline="", encoding="utf-8-sig") as handle:
            header = next(csv.reader(handle), [])
        if len(set(header)) != len(header):
            raise ValueError("CSV column names must be unique.")
        frame = pd.read_csv(file_path, dtype=object, nrows=MAX_ROWS + 1, encoding="utf-8-sig", on_bad_lines="error")
    elif extension == "json":
        with open(file_path, encoding="utf-8-sig") as handle:
            values = json.load(handle)
        if isinstance(values, dict) and isinstance(values.get("records"), list):
            values = values["records"]
        if not isinstance(values, list) or any(not isinstance(v, dict) for v in values):
            raise ValueError("JSON must contain an array of flat records or a records array.")
        if len(values) > MAX_ROWS:
            raise ValueError("Decoded input exceeds one million rows.")
        frame = pd.DataFrame(values)
    elif extension == "xlsx":
        with zipfile.ZipFile(file_path) as archive:
            members = archive.infolist()
            if len(members) > 20000 or sum(v.file_size for v in members) > MAX_DECODED_BYTES:
                raise ValueError("Spreadsheet decompression exceeds the 200 MiB budget.")
            if any("vbaProject" in v.filename or v.filename.endswith(".bin") for v in members):
                raise ValueError("Macro-enabled spreadsheets are not supported.")
        from openpyxl import load_workbook
        book = load_workbook(file_path, read_only=True, data_only=False, keep_links=False)
        sheet = book.worksheets[0]
        if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
            raise ValueError("Spreadsheet exceeds one million rows.")
        if sheet.max_column and sheet.max_column > MAX_COLUMNS:
            raise ValueError("Spreadsheet exceeds 128 columns.")
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            raise ValueError("Spreadsheet is empty.")
        rows = []
        for i, row in enumerate(iterator):
            if i >= MAX_ROWS:
                raise ValueError("Spreadsheet exceeds one million rows.")
            rows.append(row)
        frame = pd.DataFrame(rows, columns=header)
        book.close()
    elif extension in ("parquet", "canonical"):
        import pyarrow.parquet as pq
        info = pq.ParquetFile(file_path)
        if info.metadata.num_rows > MAX_ROWS or len(info.schema_arrow) > MAX_COLUMNS:
            raise ValueError("Parquet dimensions exceed the ingestion budget.")
        if sum(info.metadata.row_group(i).total_byte_size for i in range(info.metadata.num_row_groups)) > MAX_DECODED_BYTES:
            raise ValueError("Parquet decoded size exceeds 200 MiB.")
        frame = info.read().to_pandas()
    else:
        raise ValueError("Supported formats are CSV, XLSX, JSON and Parquet.")
    return validate_frame(frame)


def missing_mask(series):
    return series.isna() | series.map(lambda v: isinstance(v, str) and not v.strip())


def profile(frame):
    fields, notices = [], []
    duplicate_count = int(frame.duplicated().sum())
    if duplicate_count:
        notices.append(f"{duplicate_count:,} duplicate rows detected. Review before removing them.")
    total_issues = duplicate_count
    currency_values = set()
    for name in frame.columns:
        values = frame[name]
        missing = missing_mask(values)
        present = values[~missing]
        text = present.astype(str)
        numeric = pd.to_numeric(present, errors="coerce")
        numeric = numeric.where(np.isfinite(numeric), np.nan)
        is_numeric = len(present) > 0 and numeric.notna().mean() >= .8
        kind = "number" if is_numeric else "text"
        invalid = int(numeric.isna().sum()) if is_numeric else 0
        if len(present) and text.str.lower().isin(["true", "false"]).all():
            kind = "boolean"
        elif len(present) and text.str.match(r"^\d{4}-\d{2}-\d{2}(?:[T ].*)?$").mean() >= .8:
            kind = "date"
            invalid = int(pd.to_datetime(present, errors="coerce", utc=True, format="mixed").isna().sum())
        ambiguous = int(text.str.match(r"^(?:0?[1-9]|1[0-2])/(?:0?[1-9]|1[0-2])/\d{4}$").sum())
        formula = int(text.str.match(r"^(?:[=@]|[+-][A-Za-z(])").sum())
        if ambiguous:
            notices.append(f"{name}: {ambiguous:,} ambiguous day/month dates need an explicit format.")
        if formula:
            notices.append(f"{name}: {formula:,} formula-like cells are preserved as text and never executed.")
        if is_numeric and invalid:
            notices.append(f"{name}: {invalid:,} values do not match inferred numeric type.")
        if re.search(r"(?:currency|ccy)", name, re.I):
            currency_values.update(text.str.upper().unique().tolist())
        for symbol, code in [("$", "USD"), ("€", "EUR"), ("£", "GBP")]:
            if text.str.contains(re.escape(symbol), regex=True).any():
                currency_values.add(code)
        if re.search(r"(?:size|weight|distance|volume|duration)", name, re.I):
            units = set(text.str.extract(r"\d\s*([a-zA-Z]+)$", expand=False).dropna().str.lower())
            if len(units) > 1:
                notices.append(f"{name}: mixed units ({', '.join(sorted(units)[:6])}); conversion is not automatic.")
        total_issues += int(missing.sum()) + invalid + ambiguous + formula
        fields.append({"name": name, "type": kind, "missing": int(missing.sum()), "unique": int(present.nunique()), "invalid": invalid, "examples": clean(present.head(3).tolist())})
    if len(currency_values) > 1:
        notices.append(f"Mixed currencies detected ({', '.join(sorted(currency_values)[:10])}); totals require an explicit conversion policy.")
        total_issues += 1
    return {"rows": len(frame), "columns": len(frame.columns), "fields": fields, "preview": clean(frame.head(50).to_dict("records")), "warnings": notices, "issues": total_issues, "duplicates": duplicate_count}


def apply_transforms(frame, operations):
    if not isinstance(operations, list) or len(operations) > 20:
        raise ValueError("A recipe must contain at most 20 transformations.")
    frame = frame.copy()
    for operation in operations:
        if not isinstance(operation, dict):
            raise ValueError("Invalid transformation.")
        kind = operation.get("type")
        column = operation.get("column")
        if kind not in ("drop_duplicates", "select_columns") and column not in frame.columns:
            raise ValueError("Transformation references an unknown column.")
        if kind == "drop_duplicates":
            frame = frame.drop_duplicates()
        elif kind == "trim":
            frame[column] = frame[column].map(lambda v: v.strip() if isinstance(v, str) else v)
        elif kind == "drop_missing":
            frame = frame[~missing_mask(frame[column])]
        elif kind == "fill_missing":
            value = operation.get("value")
            if isinstance(value, (dict, list)):
                raise ValueError("Replacement values must be scalar.")
            frame.loc[missing_mask(frame[column]), column] = value
        elif kind == "rename":
            target = operation.get("name")
            if not isinstance(target, str) or not target.strip() or len(target) > 120 or target in frame.columns:
                raise ValueError("New column name must be unique and nonempty.")
            frame = frame.rename(columns={column: target})
        elif kind == "select_columns":
            columns = operation.get("columns")
            if not isinstance(columns, list) or not columns or any(c not in frame.columns for c in columns):
                raise ValueError("Select an existing set of columns.")
            frame = frame[columns]
        elif kind == "cast":
            target = operation.get("to")
            before = ~missing_mask(frame[column])
            if target == "number":
                converted = pd.to_numeric(frame[column], errors="coerce")
                converted = converted.where(np.isfinite(converted), np.nan)
            elif target == "date":
                mode = operation.get("format", "iso")
                formats = {"iso": "ISO8601", "day_first": "%d/%m/%Y", "month_first": "%m/%d/%Y"}
                if mode not in formats:
                    raise ValueError("Choose iso, day_first or month_first date format.")
                converted = pd.to_datetime(frame[column], errors="coerce", utc=True, format=formats[mode])
            elif target == "boolean":
                converted = frame[column].astype(str).str.lower().map({"true": True, "false": False, "1": True, "0": False})
            elif target == "text":
                converted = frame[column].map(lambda v: None if pd.isna(v) else str(v))
            else:
                raise ValueError("Supported types are number, date, boolean and text.")
            invalid = int((before & converted.isna()).sum())
            if invalid and operation.get("invalid") != "null":
                raise ValueError(f"{invalid} values cannot be converted. Fix the source or explicitly choose invalid=null.")
            frame[column] = converted
        else:
            raise ValueError("Unsupported transformation. Arbitrary expressions are not allowed.")
    return validate_frame(frame.reset_index(drop=True))


def canonical_write(frame, destination):
    # Arrow needs consistent physical types. Mixed object values stay literal strings.
    canonical = frame.copy()
    for column in canonical.columns:
        if canonical[column].dtype == object:
            canonical[column] = canonical[column].map(lambda v: None if pd.isna(v) else str(v))
    canonical.to_parquet(destination, index=False, engine="pyarrow", compression="zstd")
    os.chmod(destination, 0o600)
    with open(destination, "rb") as handle:
        digest = hashlib.file_digest(handle, "sha256").hexdigest()
    return digest


def base_result(kind, title, description, columns, rows, statistics=None, warnings_list=None):
    return {"kind": kind, "title": title, "description": description, "columns": columns, "rows": clean(rows), "statistics": clean(statistics or {}), "warnings": warnings_list or []}


def bootstrap_mean(values, rng, repetitions=1000):
    values = np.asarray(values, dtype=float)
    if len(values) < 2:
        return None, None
    # Limit each allocation to one bounded resample; never allocate repetitions x observations.
    means = np.array([rng.choice(values, len(values), replace=True).mean() for _ in range(repetitions)])
    return float(np.quantile(means, .025)), float(np.quantile(means, .975))


def bh_adjust(p_values):
    p = np.asarray(p_values, dtype=float)
    order = np.argsort(p)
    ranked = p[order] * len(p) / np.arange(1, len(p) + 1)
    ranked = np.minimum.accumulate(ranked[::-1])[::-1]
    output = np.empty(len(p))
    output[order] = np.minimum(1, ranked)
    return output.tolist()


def lab(data):
    kind = data.get("kind", "forecast")
    series = pd.DataFrame(data.get("series", []))
    options = data.get("options", {})
    if not isinstance(options, dict):
        raise ValueError("Analysis options must be an object.")
    if kind in ("forecast", "anomalies", "scenario"):
        if series.empty or "revenue" not in series or "month" not in series:
            return base_result(kind, "No observations", "No monthly observations match these filters.", [], [], warnings_list=["Change the filters or publish a dataset."])
        series = series.sort_values("month")
        y = np.asarray(series["revenue"], dtype=float)
        if not np.all(np.isfinite(y)):
            raise ValueError("Analysis requires finite observations; missing data must be resolved explicitly.")
        months = pd.to_datetime(series["month"], utc=True).dt.tz_localize(None).dt.to_period("M")
        if len(months) > 1 and not np.all(np.diff(months.astype('int64')) == 1):
            raise ValueError("Time series has missing months. Choose a contiguous interval; missing values are not treated as zero.")
    if kind == "anomalies":
        if len(y) < 24:
            return base_result(kind, "More history needed", "Seasonal residual analysis requires 24 monthly observations.", [], [], {"observations": len(y)}, ["The system does not fit seasonal models on insufficient history."])
        from statsmodels.tsa.seasonal import STL
        fit = STL(y, period=12, robust=True).fit()
        median = float(np.median(fit.resid))
        mad = float(np.median(np.abs(fit.resid - median)))
        scale = max(mad * 1.4826, float(np.std(fit.resid)) * .1, 1e-8)
        scores = (fit.resid - median) / scale
        # With only a few annual cycles STL can absorb a single spike into seasonality.
        # An independent Theil-Sen trend + monthly-median check protects that failure mode.
        from scipy.stats import theilslopes
        slope, intercept, _, _ = theilslopes(y, np.arange(len(y)))
        trend = intercept + slope * np.arange(len(y))
        detrended = y - trend
        expected_guard = trend + np.array([np.median(detrended[i % 12::12]) for i in range(len(y))])
        guard_residual = y - expected_guard
        guard_center = float(np.median(guard_residual))
        guard_scale = max(float(np.median(np.abs(guard_residual - guard_center))) * 1.4826, float(np.std(guard_residual)) * .1, 1e-8)
        guard_scores = (guard_residual - guard_center) / guard_scale
        use_guard = np.abs(guard_scores) > np.abs(scores)
        scores = np.where(use_guard, guard_scores, scores)
        expected = np.where(use_guard, expected_guard, fit.trend + fit.seasonal)
        residual = y - expected
        rows = [{"month": str(row["month"]), "revenue": float(y[i]), "expected": float(expected[i]), "residual": float(residual[i]), "score": float(scores[i]), "anomaly": bool(abs(scores[i]) > 3.5)} for i, (_, row) in enumerate(series.iterrows())]
        return base_result(kind, "Seasonal anomaly inspection", "Robust STL plus an independent Theil-Sen/monthly-median guard; either scaled-MAD residual score exceeding 3.5 flags an observation.", ["month", "revenue", "expected", "residual", "score", "anomaly"], rows, {"observations": len(y), "flagged": int(sum(abs(scores) > 3.5)), "threshold": 3.5}, ["Anomalies are statistical exceptions, not evidence of fraud or causation."])
    if kind == "forecast":
        if len(y) < 30:
            return base_result(kind, "More history needed", "Forecast validation needs 30 months: 24 for annual seasonality and 6 holdout origins.", [], [], {"observations": len(y)}, ["No forecast was fitted."])
        from statsmodels.tsa.holtwinters import ExponentialSmoothing
        horizon = int(options.get("horizon", 3))
        if not 1 <= horizon <= 6:
            raise ValueError("Forecast horizon must be between one and six months.")
        errors = {"seasonal_naive": [], "exponential_smoothing": []}
        actuals, predictions = [], {k: [] for k in errors}
        for stop in range(len(y) - 6, len(y)):
            train = y[:stop]
            actuals.append(float(y[stop]))
            predictions["seasonal_naive"].append(float(train[-12]))
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                fitted = ExponentialSmoothing(train, trend="add", damped_trend=True, seasonal="add", seasonal_periods=12, initialization_method="estimated").fit(optimized=True)
            predictions["exponential_smoothing"].append(float(fitted.forecast(1)[0]))
        for method in errors:
            errors[method] = np.array(actuals) - np.array(predictions[method])
        selected = min(errors, key=lambda m: float(np.mean(np.abs(errors[m]))))
        if selected == "seasonal_naive":
            forecast = np.array([y[-12 + (i % 12)] for i in range(horizon)])
        else:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                fitted = ExponentialSmoothing(y, trend="add", damped_trend=True, seasonal="add", seasonal_periods=12, initialization_method="estimated").fit(optimized=True)
            forecast = fitted.forecast(horizon)
        residuals = np.asarray(errors[selected])
        scale = max(float(np.sqrt(np.mean(residuals ** 2))), 0)
        last = pd.Timestamp(series.iloc[-1]["month"])
        rows = []
        for i, point in enumerate(forecast, 1):
            spread = 1.96 * scale * math.sqrt(i)
            rows.append({"month": (last + pd.DateOffset(months=i)).strftime("%Y-%m"), "forecast": float(point), "lower": float(point - spread), "upper": float(point + spread)})
        previous_scale = np.sqrt(np.maximum(np.cumsum(residuals[:-1] ** 2) / np.arange(1, len(residuals)), 0))
        coverage = float(np.mean(np.abs(residuals[1:]) <= 1.96 * previous_scale)) if len(residuals) > 1 else None
        denom = np.abs(actuals) + np.abs(predictions[selected])
        smape = float(np.mean(np.divide(2 * np.abs(residuals), denom, out=np.zeros_like(residuals), where=denom != 0)) * 100)
        return base_result(kind, "Revenue forecast with holdout evidence", "Select seasonal-naive or additive damped exponential smoothing by six rolling one-step holdouts. Intervals use holdout RMSE and square-root horizon scaling.", ["month", "forecast", "lower", "upper"], rows, {"model": selected, "trainingMonths": len(y), "holdoutOrigins": 6, "mae": float(np.mean(np.abs(residuals))), "sMAPE": smape, "naiveMAE": float(np.mean(np.abs(errors["seasonal_naive"]))), "etsMAE": float(np.mean(np.abs(errors["exponential_smoothing"]))), "backtestCoverage": coverage, "intervalLevel": .95}, ["Intervals are approximate; six holdouts provide limited coverage evidence. Future structural changes are not modeled."])
    if kind == "scenario":
        growth = float(options.get("growth", 10))
        reduction = float(options.get("costReduction", 15))
        if not (-90 <= growth <= 200 and -100 <= reduction <= 100):
            raise ValueError("Growth must be -90 to 200%; cost reduction must be -100 to 100%.")
        revenue = float(y.sum())
        cost = float(pd.to_numeric(series["cost"]).sum())
        projected_revenue, projected_cost = revenue * (1 + growth / 100), cost * (1 - reduction / 100)
        return base_result(kind, "Assumption-driven scenario", "Apply explicit revenue growth and infrastructure cost assumptions to the selected period.", ["scenario", "revenue", "cost", "margin"], [{"scenario": "Baseline", "revenue": revenue, "cost": cost, "margin": revenue - cost}, {"scenario": "Scenario", "revenue": projected_revenue, "cost": projected_cost, "margin": projected_revenue - projected_cost}], {"growthPercent": growth, "costReductionPercent": reduction, "marginChange": projected_revenue - projected_cost - revenue + cost}, ["This is deterministic arithmetic, not a causal forecast or guaranteed savings."])
    if kind == "cohorts":
        frame = pd.DataFrame(data.get("snapshots", []))
        if frame.empty:
            return base_result(kind, "Customer cohorts", "Monthly paid-customer retention.", [], [])
        frame["month"] = pd.to_datetime(frame["month"]).dt.to_period("M")
        frame["mrr"] = pd.to_numeric(frame["mrr"])
        first = frame[frame["mrr"] > 0].groupby("customer_id")["month"].min()
        frame = frame[frame["customer_id"].isin(first.index)].copy()
        frame["cohort"] = frame["customer_id"].map(first)
        frame = frame[frame["month"] >= frame["cohort"]]
        rows = []
        for cohort, cohort_rows in frame.groupby("cohort", sort=True):
            opening = cohort_rows[cohort_rows["month"] == cohort]
            customers = opening["customer_id"].nunique()
            opening_mrr = float(opening["mrr"].sum())
            for month, bucket in cohort_rows.groupby("month", sort=True):
                rows.append({"cohort": str(cohort), "month": str(month), "age": month.ordinal - cohort.ordinal, "startingCustomers": int(customers), "activeCustomers": int(bucket[bucket["mrr"] > 0]["customer_id"].nunique()), "retention": float(bucket[bucket["mrr"] > 0]["customer_id"].nunique() / customers * 100) if customers else None, "revenueRetention": float(bucket["mrr"].sum() / opening_mrr * 100) if opening_mrr else None})
        return base_result(kind, "Paid-customer cohorts", "Group customers by first paid subscription month; track active logos and recurring revenue for the original cohort.", ["cohort", "month", "age", "startingCustomers", "activeCustomers", "retention", "revenueRetention"], rows, {"cohorts": len(first.unique()), "customers": len(first)}, ["Cohorts use observed subscription history; history before the dataset boundary is unavailable."])
    if kind == "comparison":
        from scipy.stats import mannwhitneyu
        values = pd.DataFrame(data.get("customers", []))
        if values.empty:
            return base_result(kind, "Customer-level comparison", "No customer observations match these filters.", [], [])
        groups = list(values.groupby("region", sort=True))
        if len(groups) < 2:
            return base_result(kind, "Another segment is needed", "Comparison requires at least two regions.", [], [], warnings_list=["Clear the region filter to compare segments."])
        rng = np.random.default_rng(7139)
        baseline_name, baseline_frame = groups[0]
        baseline = pd.to_numeric(baseline_frame["revenue"]).to_numpy(float)
        rows = []
        for name, group in groups[1:]:
            sample = pd.to_numeric(group["revenue"]).to_numpy(float)
            if len(sample) < 2 or len(baseline) < 2:
                continue
            differences = np.array([rng.choice(sample, len(sample), True).mean() - rng.choice(baseline, len(baseline), True).mean() for _ in range(1000)])
            p = float(mannwhitneyu(sample, baseline, alternative="two-sided").pvalue)
            rows.append({"segment": str(name), "baseline": str(baseline_name), "customers": len(sample), "baselineCustomers": len(baseline), "meanRevenue": float(np.mean(sample)), "meanDifference": float(np.mean(sample) - np.mean(baseline)), "lower": float(np.quantile(differences, .025)), "upper": float(np.quantile(differences, .975)), "pValue": p})
        if rows:
            for row, adjusted in zip(rows, bh_adjust([r["pValue"] for r in rows])):
                row["adjustedP"] = adjusted
        return base_result(kind, "Customer-level regional comparison", "Customer-level bootstrap intervals for mean differences; Mann-Whitney distribution tests with Benjamini-Hochberg correction.", ["segment", "baseline", "customers", "baselineCustomers", "meanRevenue", "meanDifference", "lower", "upper", "adjustedP"], rows, {"resamples": 1000, "intervalLevel": .95}, ["Distribution-test p-values do not test the mean difference. Customer mix and selection may confound these associations."])
    if kind == "funnel":
        rows = data.get("funnel", [])
        previous = None
        for row in rows:
            count = int(row["customers"])
            row["conversion"] = count / previous * 100 if previous else (100 if count else None)
            previous = count
        return base_result(kind, "Observed customer lifecycle", "Count customers whose first recorded activity proceeds to usage and then a later positive invoice within 90 days.", ["stage", "customers", "conversion"], rows, {"windowDays": 90}, ["Stages represent observed events, not website visits or attributed causal effects."])
    if kind == "segments":
        rows = data.get("segments", [])
        total = sum(float(row.get("revenue", 0)) for row in rows)
        for row in rows:
            row["contribution"] = float(row["revenue"]) / total * 100 if total else None
        return base_result(kind, "Revenue contribution by region", "Compare billed revenue and customer counts across authorized regional segments.", ["name", "revenue", "customers", "contribution"], rows, {"totalRevenue": total}, ["Contribution is descriptive; segment size and mix may explain differences."])
    raise ValueError("Unknown analysis. Choose forecast, anomalies, scenario, cohorts, comparison, funnel or segments.")


def main(payload):
    task = payload.get("task")
    if task == "query":
        import duckdb
        frame = read_table(payload["path"], "canonical")
        # SQL is compiled and AST-validated by the server. No uploaded code or SQL is accepted.
        connection = duckdb.connect(config={"enable_external_access": "false", "threads": "1", "memory_limit": "256MB"})
        try:
            connection.register("selected_dataset", frame)
            result = connection.execute(payload["sql"], payload.get("values", [])).fetchdf()
            if len(result) > 10000:
                raise ValueError("Query results exceed 10,000 rows.")
            return {"columns": result.columns.tolist(), "rows": clean(result.to_dict("records"))}
        finally:
            connection.close()
    if task == "lab":
        return lab(payload)
    if task == "rows":
        frame = read_table(payload["path"], "canonical")
        offset, limit = int(payload.get("offset", 0)), int(payload.get("limit", 100))
        if offset < 0 or not 1 <= limit <= 1000:
            raise ValueError("Invalid page range.")
        return {"rows": clean(frame.iloc[offset:offset + limit].to_dict("records")), "total": len(frame), "offset": offset, "limit": limit}
    if task not in ("profile", "transform"):
        raise ValueError("Unknown analytical task.")
    frame = read_table(payload["path"], payload.get("extension", "canonical"))
    if task == "transform":
        frame = apply_transforms(frame, payload.get("operations", []))
    result = profile(frame)
    result["hash"] = canonical_write(frame, payload["destination"])
    result["software"] = {"pandas": pd.__version__, "numpy": np.__version__, "recipeVersion": "1"}
    return result


if __name__ == "__main__":
    try:
        request = sys.stdin.read(16 * 1024 * 1024 + 1)
        if len(request) > 16 * 1024 * 1024:
            raise ValueError("Input exceeds its limit.")
        result = main(json.loads(request))
        sys.stdout.write(json.dumps(clean(result), allow_nan=False, separators=(",", ":")))
    except Exception as exc:
        sys.stdout.write(json.dumps({"error": str(exc)[:500], "code": "invalid_analytical_input"}))
        sys.exit(1)
